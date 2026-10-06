"""Unit tests for the ``read_document`` chat tool (attached-document content extraction).

Same style as ``test_document_tools.py``: a real :class:`LocalStorage` at tmp_path holds
the asset bytes, and the drive runs on fake repositories (``make_drive``) — the tool now
authorizes through ``ensure_asset_readable`` before touching storage, and reads the
request user from the ``request_user`` contextvar, exactly like a real chat turn.
"""
from __future__ import annotations

import hashlib
from io import BytesIO
from uuid import uuid4

from agent import Context
from agent.engine.decisions import ToolExecution
from agent.engine.runtime import ToolRuntime
from core.infrastructure.storage import LocalStorage, object_key

from apps.api.tools import read_document_tool
from tests._drive_fakes import make_drive


def _ctx(tmp_path) -> Context:
    ctx = Context()
    ctx.provide("storage", LocalStorage(tmp_path))
    ctx.provide("session_factory", object)  # unused: the drive is monkeypatched
    return ctx


async def _register(monkeypatch, tmp_path, name: str, data: str | bytes,
                    user=None) -> tuple[ToolRuntime, str]:
    """Store ``data`` as ``user``'s asset named ``name``; return (runtime, asset_id).

    Also binds ``request_user`` to ``user`` so the tool's ACL check passes — same context
    the chat router establishes around a turn.
    """
    if isinstance(data, str):
        data = data.encode()
    storage = LocalStorage(tmp_path)
    sha = hashlib.sha256(data).hexdigest()
    await storage.put(object_key(sha), data)
    drive = make_drive(tmp_path)
    uid = user or uuid4()
    asset = await drive.assets.create(uid, name, object_sha256=sha)
    monkeypatch.setattr(read_document_tool, "DriveService", lambda sf: drive)
    read_document_tool.request_user.set(uid)
    runtime = ToolRuntime()
    read_document_tool.register(runtime, _ctx(tmp_path), llm=None)
    return runtime, str(asset.id)


async def _call(runtime: ToolRuntime, asset_id: str, pages=None):
    args = {"asset_id": asset_id}
    if pages is not None:
        args["pages"] = pages
    return await runtime.execute(ToolExecution("c1", "read_document", args))


def _err(res) -> str:
    """Error text of a failed tool result (``ToolExecutionFailure`` carries no ``value``)."""
    return str(res.error.message) + " " + " ".join(str(getattr(b, "text", b)) for b in res.content)


async def test_reads_plain_text(monkeypatch, tmp_path):
    runtime, asset_id = await _register(monkeypatch, tmp_path, "notes.txt", b"hello document")
    res = await _call(runtime, asset_id)
    assert res.is_error is False
    assert "hello document" in str(res.value)


async def test_reads_markdown_and_csv(monkeypatch, tmp_path):
    runtime, a1 = await _register(monkeypatch, tmp_path, "doc.md", b"# Heading\nbody")
    assert "Heading" in str((await _call(runtime, a1)).value)
    runtime, a2 = await _register(monkeypatch, tmp_path, "t.csv", b"a,b\n1,2")
    assert "1,2" in str((await _call(runtime, a2)).value)


async def test_reads_docx(monkeypatch, tmp_path):
    import docx

    buf = BytesIO()
    d = docx.Document()
    d.add_paragraph("docx paragraph one")
    d.save(buf)
    runtime, asset_id = await _register(monkeypatch, tmp_path, "report.docx", buf.getvalue())
    res = await _call(runtime, asset_id)
    assert res.is_error is False
    assert "docx paragraph one" in str(res.value)


async def test_reads_xlsx(monkeypatch, tmp_path):
    from openpyxl import Workbook

    buf = BytesIO()
    wb = Workbook()
    wb.active["A1"] = "quarter"
    wb.active["B1"] = 42
    wb.save(buf)
    runtime, asset_id = await _register(monkeypatch, tmp_path, "book.xlsx", buf.getvalue())
    res = await _call(runtime, asset_id)
    assert res.is_error is False
    assert "quarter" in str(res.value)
    assert "42" in str(res.value)


async def test_reads_pdf(monkeypatch, tmp_path):
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "pdf body text")
    data = doc.tobytes()
    doc.close()
    runtime, asset_id = await _register(monkeypatch, tmp_path, "paper.pdf", data)
    res = await _call(runtime, asset_id)
    assert res.is_error is False
    assert "pdf body text" in str(res.value)


async def test_reads_pptx_with_notes(monkeypatch, tmp_path):
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "DCS overview slide"
    slide.notes_slide.notes_text_frame.text = "speaker note: fieldbus protocol"
    buf = BytesIO()
    prs.save(buf)
    runtime, asset_id = await _register(monkeypatch, tmp_path, "deck.pptx", buf.getvalue())
    res = await _call(runtime, asset_id)
    assert res.is_error is False
    assert "## slide 1" in str(res.value)
    assert "DCS overview slide" in str(res.value)
    assert "speaker note: fieldbus protocol" in str(res.value)


def _as_alt_package(raw: bytes, alt_main: str) -> bytes:
    """Rewrite a .pptx package's main content type (template/slideshow), as PowerPoint saves do."""
    import zipfile

    src = zipfile.ZipFile(BytesIO(raw))
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.namelist():
            data = src.read(item)
            if item == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.presentationml"
                    b".presentation.main+xml",
                    (
                        b"application/vnd.openxmlformats-officedocument.presentationml."
                        + alt_main.encode()
                        + b".main+xml"
                    ),
                )
            dst.writestr(item, data)
    return out.getvalue()


async def test_reads_potx_template(monkeypatch, tmp_path):
    """The reported bug: a lecture .potx attach must extract, not fail and let the model
    summarize an earlier document."""
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "template deck title"
    buf = BytesIO()
    prs.save(buf)
    potx = _as_alt_package(buf.getvalue(), "template")
    runtime, asset_id = await _register(monkeypatch, tmp_path, "lecture.potx", potx)
    res = await _call(runtime, asset_id)
    assert res.is_error is False
    assert "template deck title" in str(res.value)


async def test_unsupported_message_forbids_stale_substitution(monkeypatch, tmp_path):
    runtime, asset_id = await _register(monkeypatch, tmp_path, "legacy.ppt", b"\xd0\xcf\x11\xe0junk")
    res = await _call(runtime, asset_id)
    assert "Cannot extract text" in str(res.value)
    assert "must NOT substitute" in str(res.value)


async def test_image_routed_to_vision_hint(monkeypatch, tmp_path):
    png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 32
    runtime, asset_id = await _register(monkeypatch, tmp_path, "shot.png", png)
    res = await _call(runtime, asset_id)
    assert res.is_error is False
    assert "vision" in str(res.value)


async def test_unsupported_type_returns_message(monkeypatch, tmp_path):
    runtime, asset_id = await _register(monkeypatch, tmp_path, "pkg.zip", b"PK\x03\x04junk")
    res = await _call(runtime, asset_id)
    assert res.is_error is False
    assert "Cannot extract text" in str(res.value)


async def test_legacy_xls_rejected_with_resave_hint(monkeypatch, tmp_path):
    runtime, asset_id = await _register(monkeypatch, tmp_path, "old.xls", b"\xd0\xcf\x11\xe0junk")
    res = await _call(runtime, asset_id)
    assert "xlsx" in str(res.value)


async def _register_doc(monkeypatch, tmp_path, images, skipped=0):
    """.doc attach with stubbed text extraction + scanner; returns (runtime, id, drive)."""
    data = b"\xd0\xcf\x11\xe0doc container"
    sha = hashlib.sha256(data).hexdigest()
    await LocalStorage(tmp_path).put(object_key(sha), data)
    drive = make_drive(tmp_path)
    uid = uuid4()
    asset = await drive.assets.create(uid, "report.doc", object_sha256=sha)
    monkeypatch.setattr(read_document_tool, "DriveService", lambda sf: drive)
    read_document_tool.request_user.set(uid)

    async def _text(_data, _name, _llm):
        return "quarterly numbers [pic] and more"

    monkeypatch.setattr(read_document_tool, "extract_document_text", _text)
    monkeypatch.setattr(read_document_tool, "scan_doc_images", lambda _d: (images, skipped))

    runtime = ToolRuntime()
    read_document_tool.register(runtime, _ctx(tmp_path), llm=None)
    return runtime, str(asset.id), drive


async def test_doc_images_minted_and_asset_ids_listed(monkeypatch, tmp_path):
    img = {"name": "doc_1.png", "mime": "image/png", "data": b"PNGBYTES"}
    runtime, asset_id, drive = await _register_doc(monkeypatch, tmp_path, [img])
    res = await _call(runtime, asset_id)
    out = str(res.value)
    assert res.is_error is False
    assert "quarterly numbers" in out
    assert "embedded image" in out and "vision" in out

    # exactly one derived asset, bound to the source doc, in the RAG folder
    saved = [a for a in drive.assets.rows.values() if str(a.id) != asset_id]
    assert len(saved) == 1
    assert saved[0].source_asset_id is not None
    assert saved[0].folder_path == "RAG 图片/report"
    assert f"asset_id {saved[0].id}" in out


async def test_doc_image_mint_dedupes_existing_asset(monkeypatch, tmp_path):
    img = {"name": "doc_1.png", "mime": "image/png", "data": b"PNGBYTES"}
    runtime, asset_id, drive = await _register_doc(monkeypatch, tmp_path, [img])
    await _call(runtime, asset_id)  # first call mints
    res2 = await _call(runtime, asset_id)  # second reuses the same asset row
    derived = [a for a in drive.assets.rows.values() if a.source_asset_id is not None]
    assert len(derived) == 1
    assert hashlib.sha256(img["data"]).hexdigest() == derived[0].object_sha256
    assert f"asset_id {derived[0].id}" in str(res2.value)


async def test_doc_without_images_unaffected(monkeypatch, tmp_path):
    runtime, asset_id, _ = await _register_doc(monkeypatch, tmp_path, [], skipped=2)
    out = str((await _call(runtime, asset_id)).value)
    assert "metafile" in out and "could not be rendered" in out


async def test_long_document_is_truncated(monkeypatch, tmp_path):
    from apps.api.tools.read_document_tool import MAX_OUTPUT_CHARS

    payload = b"x" * (MAX_OUTPUT_CHARS + 500)
    runtime, asset_id = await _register(monkeypatch, tmp_path, "big.txt", payload)
    res = await _call(runtime, asset_id)
    assert "truncated" in str(res.value)
    assert len(str(res.value)) < MAX_OUTPUT_CHARS + 200


async def test_missing_asset_raises(monkeypatch, tmp_path):
    runtime, _ = await _register(monkeypatch, tmp_path, "a.txt", b"x")
    res = await _call(runtime, str(uuid4()))
    assert res.is_error is True


# ── pages: parser matrix ──

def test_parse_pages_spec_dedupes_and_sorts():
    from core.infrastructure.pdf_pages import parse_pages_spec

    assert parse_pages_spec("1,2-4,4") == [1, 2, 3, 4]
    assert parse_pages_spec(" 1, 3-5 , 8 ") == [1, 3, 4, 5, 8]
    assert parse_pages_spec("2-2") == [2]
    assert parse_pages_spec("3") == [3]


def test_parse_pages_spec_rejects_malformed():
    import pytest
    from core.infrastructure.pdf_pages import parse_pages_spec

    for bad in ["", "   ", "1-", "-3", "abc", "1,,3", "5-2", "0", "2-0", "1,0"]:
        with pytest.raises(ValueError):
            parse_pages_spec(bad)


def test_parse_pages_spec_span_checked_before_expansion():
    """A pathological range must raise WITHOUT building the expanded list."""
    import pytest
    from core.infrastructure.pdf_pages import parse_pages_spec

    with pytest.raises(ValueError, match="spans"):
        parse_pages_spec("1-1000000000")


def test_parse_pages_spec_incremental_cap():
    """16 unique pages pass; the 17th raises as it is added, never after."""
    import pytest
    from core.infrastructure.pdf_pages import parse_pages_spec

    assert len(parse_pages_spec("1-16")) == 16
    with pytest.raises(ValueError, match="per-call limit"):
        parse_pages_spec("1-16,17")


# ── pages: execute-path contract ──

def _pdf_pages(texts: list[str]) -> bytes:
    import fitz

    doc = fitz.open()
    for t in texts:
        page = doc.new_page()
        if t:
            page.insert_text((72, 72), t)
    data = doc.tobytes()
    doc.close()
    return data


async def test_pages_empty_string_and_wrong_type_are_errors(monkeypatch, tmp_path):
    runtime, asset_id = await _register(monkeypatch, tmp_path, "p.pdf", _pdf_pages(["a"]))
    res = await _call(runtime, asset_id, pages="")
    assert res.is_error is True and "non-empty" in _err(res)
    res = await _call(runtime, asset_id, pages=3)
    # schema validation rejects non-string pages before the executor ever runs
    assert res.is_error is True and "not of type 'string'" in _err(res)


async def test_pdf_page_scoped_reads_only_requested_pages(monkeypatch, tmp_path):
    data = _pdf_pages(["page one text", "page two text", "page three text"])
    runtime, asset_id = await _register(monkeypatch, tmp_path, "paper.pdf", data)
    res = await _call(runtime, asset_id, pages="1,3")
    out = str(res.value)
    assert res.is_error is False
    assert "page one text" in out and "page three text" in out
    assert "page two text" not in out
    assert "[[Page 1]]" in out and "[[Page 3]]" in out


async def test_pdf_out_of_range_fails_whole_call_before_extraction(monkeypatch, tmp_path):
    data = _pdf_pages(["solo page"])
    runtime, asset_id = await _register(monkeypatch, tmp_path, "paper.pdf", data)
    res = await _call(runtime, asset_id, pages="1,5")
    assert res.is_error is True
    # the valid page must NOT leak into the error, and the real count is reported
    assert "has 1 page" in _err(res) and "out of range" in _err(res)
    assert "[[Page 1]]" not in _err(res)


async def test_pdf_empty_page_is_inline_warning_not_failure(monkeypatch, tmp_path):
    data = _pdf_pages(["has text", ""])
    runtime, asset_id = await _register(monkeypatch, tmp_path, "mixed.pdf", data)
    res = await _call(runtime, asset_id, pages="1,2")
    out = str(res.value)
    assert res.is_error is False
    assert "has text" in out
    assert "No extractable text found on this page" in out


async def test_pptx_slide_scoped_and_notes_omitted(monkeypatch, tmp_path):
    from io import BytesIO

    from pptx import Presentation

    prs = Presentation()
    for i in (1, 2, 3):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = f"slide body {i}"
        slide.notes_slide.notes_text_frame.text = f"secret note {i}"
    buf = BytesIO()
    prs.save(buf)
    runtime, asset_id = await _register(monkeypatch, tmp_path, "deck.pptx", buf.getvalue())
    res = await _call(runtime, asset_id, pages="2")
    out = str(res.value)
    assert res.is_error is False
    assert "slide body 2" in out
    assert "slide body 1" not in out and "slide body 3" not in out
    assert "secret note" not in out  # the page-scoped path never parses notes

    res = await _call(runtime, asset_id, pages="9")
    assert res.is_error is True and "has 3 slide" in _err(res)


async def test_pages_rejected_for_no_page_axis(monkeypatch, tmp_path):
    runtime, docx_id = await _register(monkeypatch, tmp_path, "notes.docx", b"whatever")
    res = await _call(runtime, docx_id, pages="2")
    assert res.is_error is True
    assert "no page axis" in _err(res)
    assert "PPTX/POTX/PPSX" in _err(res)

    runtime, txt_id = await _register(monkeypatch, tmp_path, "plain.txt", b"hello")
    res = await _call(runtime, txt_id, pages="2")
    assert res.is_error is True and "no page axis" in _err(res)


async def test_pages_truncation_applies_once_to_merged_output(monkeypatch, tmp_path):
    from apps.api.tools import read_document_tool as rd

    data = _pdf_pages(["alpha " * 100, "beta " * 100])
    runtime, asset_id = await _register(monkeypatch, tmp_path, "paper.pdf", data)
    monkeypatch.setattr(rd, "MAX_OUTPUT_CHARS", 200)
    res = await _call(runtime, asset_id, pages="1,2")
    out = str(res.value)
    assert res.is_error is False
    assert "truncated" in out
    # one cap on the MERGED result: never two per-page truncation notes
    assert out.count("truncated") == 1
    assert len(out) < 200 + 200


# ── ACL before storage ──

async def test_no_request_user_fails_before_storage(monkeypatch, tmp_path):
    runtime, asset_id = await _register(monkeypatch, tmp_path, "a.txt", b"hello")
    token = read_document_tool.request_user.set(None)
    try:
        res = await _call(runtime, asset_id)
    finally:
        read_document_tool.request_user.reset(token)
    assert res.is_error is True
    assert "authenticated request user" in _err(res)


async def test_foreign_asset_denied_before_bytes_fetch(monkeypatch, tmp_path):
    runtime, asset_id = await _register(monkeypatch, tmp_path, "a.txt", b"hello")
    # attacker context: the object bytes are gone from storage — an ACL hole would surface
    # as "object bytes missing"; the ACL must fire first with a not-readable error.
    sha = hashlib.sha256(b"hello").hexdigest()
    await LocalStorage(tmp_path).delete(object_key(sha))
    read_document_tool.request_user.set(uuid4())
    res = await _call(runtime, asset_id)
    assert res.is_error is True
    assert "not readable" in _err(res)
    assert "bytes missing" not in _err(res)
