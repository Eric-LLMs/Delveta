"""Authorized document reading — the ONE asset→text path.

Both ``read_document`` (agent-facing, output-capped for the context window) and
``summary_gen`` (full text handed to the toolkit generation pipeline) read a
chat attachment / viewer document through this module, so the drive ACL check,
the storage fetch, the page-axis validation and the format extraction exist
exactly once and cannot drift between the two tools.

Nothing here decides WHAT to read: the caller supplies the trusted ``asset_id``
(resolved by the funnel from turn facts) and the optional page list (resolved by
the caller's own scope rule). This module never invents either.
"""
from __future__ import annotations

import asyncio
import logging
from dataclasses import dataclass
from pathlib import Path
from uuid import UUID

from core.application.drive_service import DriveError, DriveService
from core.infrastructure.ingest import (
    UnsupportedFileType,
    _ppt_as_presentation,
    _shape_texts,
    extract_document_text,
)
from core.infrastructure.request_context import request_user
from core.infrastructure.storage import object_key

log = logging.getLogger(__name__)

# Formats with a real page/slide axis. ``pages`` is rejected for everything else — the
# reader never substitutes a full-document read for a page-scoped request.
PAGE_AXIS_EXTS = frozenset({".pdf", ".pptx", ".potx", ".ppsx"})

# Extensions handled by the dedicated vision tool, not by text extraction.
IMAGE_EXTS = frozenset({".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".svg"})


class UnsupportedDocument(UnsupportedFileType):
    """The asset is readable, but no text extractor handles its format.

    Carries the document name so each caller can word its own user-facing error
    without re-reading the asset.
    """

    def __init__(self, name: str, reason: str) -> None:
        super().__init__(reason)
        self.name = name
        self.reason = reason


@dataclass(frozen=True)
class AssetDocument:
    """One readable drive asset: its identity, metadata, bytes and extracted text.

    ``text`` is the FULL text of the requested scope (whole document or exactly the
    requested pages/slides); the caller owns any context-window truncation.
    """

    asset_id: str
    name: str
    ext: str
    mime: str
    data: bytes
    text: str
    asset: object

    @property
    def is_image(self) -> bool:
        return self.ext in IMAGE_EXTS or self.mime.startswith("image/")

    @property
    def has_page_axis(self) -> bool:
        return self.ext in PAGE_AXIS_EXTS


def _extract_pdf_pages(data: bytes, name: str, pages: list[int]) -> str:
    """PyMuPDF per-page text for the REQUESTED pages only — never the whole document.

    No table detection, no vision-LLM transcription, no scan pass on unrequested pages.
    Bounds are validated before any page is extracted: an out-of-range request fails the
    WHOLE call (no partially-extracted valid pages leak into the error path). A requested
    page that carries no text (scanned page) is a warning inline, not a tool failure.
    """
    import pymupdf  # same dep the ingest extractor uses; import kept local & cheap

    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        count = doc.page_count
        out_of_range = [p for p in pages if p > count]
        if out_of_range:
            raise ValueError(
                f"page(s) {out_of_range} out of range: '{name}' has {count} page(s)"
            )
        parts: list[str] = []
        for page in pages:
            text = (doc[page - 1].get_text("text") or "").strip()
            if not text:
                text = (
                    "[No extractable text found on this page. PDF visual-page analysis "
                    "is outside this tool's current scope.]"
                )
            parts.append(f"[[Page {page}]]\n{text}")
        return "\n\n".join(parts)
    finally:
        doc.close()


def _extract_pptx_slides(data: bytes, name: str, pages: list[int]) -> str:
    """python-pptx per-slide text for the REQUESTED slides only (no notes parsing added —
    that stays the full-document ingest path's extra). Reuses the shape-flattening helpers
    from the core extractor so slide text is byte-identical to what the deck chunks carry.
    """
    import io

    from pptx import Presentation

    prs = Presentation(io.BytesIO(_ppt_as_presentation(data)))
    slides = list(prs.slides)
    out_of_range = [p for p in pages if p > len(slides)]
    if out_of_range:
        raise ValueError(
            f"slide(s) {out_of_range} out of range: '{name}' has {len(slides)} slide(s)"
        )
    parts: list[str] = []
    for page in pages:
        body = "\n".join(_shape_texts(slides[page - 1].shapes)).strip()
        if not body:
            body = "[No extractable text found on this slide.]"
        parts.append(f"[[Page {page}]]\n## slide {page}\n{body}")
    return "\n\n".join(parts)


async def read_asset_document(
    ctx,
    llm,
    asset_id: str,
    *,
    pages: list[int] | None = None,
    caller: str = "read_document",
) -> AssetDocument:
    """Load one drive asset the CURRENT request user is allowed to read.

    ``pages`` is an already-parsed 1-based page/slide list; ``None`` reads the whole
    document. A page list on a format without a page axis is a hard error — never a
    silent whole-document read. Images are returned with ``text=""`` and
    ``is_image=True`` so each caller can point at the ``vision`` tool.
    """
    uid = request_user.get()
    if uid is None:
        raise ValueError(
            f"{caller} requires an authenticated request user context"
        )

    # ── permission before storage: the drive ACL decides readability; bytes are only
    # fetched for assets the current request user may actually read.
    drive = DriveService(ctx.resolve("session_factory"))
    try:
        asset = await drive.ensure_asset_readable(uid, UUID(str(asset_id)))
    except DriveError as exc:
        raise ValueError(f"asset {asset_id} not readable: {exc}") from None
    if not asset.object_sha256:
        raise ValueError(f"asset {asset_id} not found or has no stored object")
    data = await ctx.resolve("storage").get(object_key(asset.object_sha256))
    if data is None:
        raise ValueError(f"object bytes missing for asset {asset_id}")

    name = getattr(asset, "name", None) or "document"
    ext = Path(name).suffix.lower()
    mime = (getattr(asset, "mime_type", None) or "").lower()
    base = AssetDocument(asset_id=str(asset_id), name=name, ext=ext, mime=mime,
                         data=data, text="", asset=asset)
    if base.is_image:
        return base

    if pages is not None:
        if ext not in PAGE_AXIS_EXTS:
            raise ValueError(
                f"'{name}' has no page axis ({ext or 'unknown format'}): pages is only "
                "supported for PDF and PowerPoint slide-based formats "
                "(PPTX/POTX/PPSX). Omit pages to read the whole document, and never "
                "substitute a full-document read for a page-scoped request."
            )
        if ext == ".pdf":
            text = await asyncio.to_thread(_extract_pdf_pages, data, name, pages)
        else:
            text = await asyncio.to_thread(_extract_pptx_slides, data, name, pages)
        return AssetDocument(asset_id=base.asset_id, name=name, ext=ext, mime=mime,
                             data=data, text=text, asset=asset)

    try:
        text = await extract_document_text(data, name, llm)
    except UnsupportedFileType as exc:
        raise UnsupportedDocument(name, str(exc)) from exc
    return AssetDocument(asset_id=base.asset_id, name=name, ext=ext, mime=mime,
                         data=data, text=text, asset=asset)
