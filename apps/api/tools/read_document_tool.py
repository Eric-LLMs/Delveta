"""``read_document``: extract the text content of a document the user attached to the chat
or has open in the viewer.

Attachments ride on the chat payload as a drive ``asset_id``; the agent only sees the
``[Attached: <name> (asset_id <id>)]`` note or the trusted Viewer Access Context stub.
Without a reader it can guess but never open
the file, so every "parse this PDF / summarize this Word doc" request fails. This tool
closes that gap: it loads the asset bytes from storage and runs the same extractor the
ingest worker uses — PDF (PyMuPDF body text, tables via the vision LLM), .docx
(python-docx), legacy .doc (antiword), Excel (.xlsx via openpyxl), PowerPoint
(.pptx/.potx/.ppsx slide text + speaker notes via python-pptx), plus plain text /
markdown / csv / json and subtitles. Images are routed to the ``vision`` tool instead,
and the output is capped so one huge document cannot flood the agent's context window.
With ``pages`` the read is page/slide-scoped: PyMuPDF/python-pptx touch ONLY the requested
pages (no whole-document pass, no table/vision work on unrequested pages), out-of-range
specs fail the whole call before extraction, formats without a page axis are rejected, and
the output cap applies once to the merged multi-page result.
Legacy .doc additionally gets an embedded-image recovery pass: antiword can only leave a
``[pic]`` placeholder, so ``scan_doc_images`` recovers the actual inline pictures, saves
them as derived drive assets (same ``RAG 图片/<doc>`` folder and dedupe the ingest worker
uses), and lists their asset_ids in the tool output so the agent can feed them to
``vision``.
"""
from __future__ import annotations

import asyncio
import hashlib
import logging
import re
from pathlib import Path
from uuid import UUID

from agent import Context, ToolExecution, ToolOutput, ToolRuntime, define_tool, text_block
from core.application.drive_service import DriveError, DriveService
from core.infrastructure.doc_images import scan_doc_images
from core.infrastructure.ingest import (
    UnsupportedFileType,
    _ppt_as_presentation,
    _shape_texts,
    extract_document_text,
)
from core.infrastructure.pdf_pages import parse_pages_spec
from core.infrastructure.request_context import request_user
from core.infrastructure.storage import object_key

log = logging.getLogger(__name__)

# One tool result must not blow up the agent's context window; the tail is dropped with a
# truncation note so the agent can tell the user it saw only the first part. It caps the
# FINAL MERGED output — page-scoped multi-page results truncate once, never per page.
MAX_OUTPUT_CHARS = 20_000

# Formats with a real page/slide axis. ``pages`` is rejected for everything else — the tool
# never substitutes a full-document read for a page-scoped request.
_PAGE_AXIS_EXTS = {".pdf", ".pptx", ".potx", ".ppsx"}

# Characters that would turn a doc title into nested folders / broken names.
_FOLDER_BAD = re.compile(r'[\\/:*?"<>|]')

# Extensions handled by the dedicated vision tool, not by text extraction.
_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".svg"}


def _extract_pdf_pages(data: bytes, name: str, pages: list[int]) -> str:
    """PyMuPDF per-page text for the REQUESTED pages only — never the whole document.

    No table detection, no vision-LLM transcription, no scan pass on unrequested pages.
    Bounds are validated before any page is extracted: an out-of-range request fails the
    WHOLE call (no partially-extracted valid pages leak into the error path). A requested
    page that carries no text (scanned page) is a warning inline, not a tool failure.
    """
    import pymupdf  # same dep the ingest extractor uses; import kept local & cheap

    doc = pymupdf.open(stream=data, filetype="pdf")
    try:
        count = doc.page_count
        out_of_range = [p for p in pages if p > count]
        if out_of_range:
            raise ValueError(
                f"page(s) {out_of_range} out of range: '{name}' has {count} page(s)"
            )
        parts: list[str] = []
        for page in pages:
            text = (doc[page - 1].get_text("text") or "").strip()
            if not text:
                text = (
                    "[No extractable text found on this page. PDF visual-page analysis "
                    "is outside this tool's current scope.]"
                )
            parts.append(f"[[Page {page}]]\n{text}")
        return "\n\n".join(parts)
    finally:
        doc.close()


def _extract_pptx_slides(data: bytes, name: str, pages: list[int]) -> str:
    """python-pptx per-slide text for the REQUESTED slides only (no notes parsing added —
    that stays the full-document ingest path's extra). Reuses the shape-flattening helpers
    from the core extractor so slide text is byte-identical to what the deck chunks carry.
    """
    import io

    from pptx import Presentation

    prs = Presentation(io.BytesIO(_ppt_as_presentation(data)))
    slides = list(prs.slides)
    out_of_range = [p for p in pages if p > len(slides)]
    if out_of_range:
        raise ValueError(
            f"slide(s) {out_of_range} out of range: '{name}' has {len(slides)} slide(s)"
        )
    parts: list[str] = []
    for page in pages:
        body = "\n".join(_shape_texts(slides[page - 1].shapes)).strip()
        if not body:
            body = "[No extractable text found on this slide.]"
        parts.append(f"[[Page {page}]]\n## slide {page}\n{body}")
    return "\n\n".join(parts)


async def _mint_doc_images(data: bytes, asset, ctx: Context) -> str:
    """Save a legacy .doc's inline pictures as drive assets; return a footer listing ids.

    Dedupes on ``(source_asset_id, sha256)`` — an image already extracted by the RAG
    worker (or a previous chat turn) resolves to the existing asset instead of a
    duplicate row. Any failure degrades to the plain text result: minting must never
    break ``read_document``.
    """
    try:
        images, skipped = await asyncio.to_thread(scan_doc_images, data)
        if not images:
            if skipped:
                return (
                    f"\n\n[{skipped} embedded metafile image(s) could not be rendered "
                    "on this server; the [pic] placeholder above is all there is.]"
                )
            return ""
        stem = _FOLDER_BAD.sub("_", Path(asset.name or "doc").stem).strip(". ") or "doc"
        folder = f"RAG 图片/{stem}"
        drive = DriveService(ctx.resolve("session_factory"))
        ids: list[str] = []
        for img in images:
            digest = hashlib.sha256(img["data"]).hexdigest()
            existing = await drive.assets.get_by_source_content(asset.id, digest)
            if existing is not None:
                ids.append(str(existing.id))
                continue
            saved = await drive.save_artifact(
                asset.user_id,
                img["name"],
                img["mime"],
                img["data"],
                folder_path=folder,
                workspace_id=asset.workspace_id,
                source_asset_id=asset.id,
            )
            ids.append(str(saved.id))
        lines = "\n".join(
            f"- doc_image {i + 1}: asset_id {aid}" for i, aid in enumerate(ids)
        )
        extra = f" ({skipped} metafile image(s) not renderable)" if skipped else ""
        return (
            f"\n\n[This .doc contains {len(ids)} embedded image(s){extra}. They are saved "
            "in the cloud drive; call the `vision` tool with an asset_id below to analyze "
            f"the picture or table it shows:\n{lines}]"
        )
    except Exception as exc:  # noqa: BLE001 — image recovery is best-effort
        log.warning("doc image minting failed for asset %s: %s", asset.id, exc)
        return ""


def register(runtime: ToolRuntime, ctx: Context, llm) -> None:
    async def read_document(args: dict, exec: ToolExecution) -> str:
        asset_id = args["asset_id"]
        # ── pages tri-state BEFORE any I/O: omitted/None → legacy full read; empty string
        # → hard error (must never be mistaken for "read everything"); non-empty → parse.
        pages_raw = args.get("pages")
        if pages_raw is not None and not isinstance(pages_raw, str):
            raise ValueError('pages must be a string like "3" or "1-3"')
        pages = parse_pages_spec(pages_raw) if pages_raw is not None else None

        # ── permission before storage: the drive ACL decides readability; bytes are only
        # fetched for assets the current request user may actually read.
        uid = request_user.get()
        if uid is None:
            raise ValueError("read_document requires an authenticated request user context")
        drive = DriveService(ctx.resolve("session_factory"))
        try:
            asset = await drive.ensure_asset_readable(uid, UUID(str(asset_id)))
        except DriveError as exc:
            raise ValueError(f"asset {asset_id} not readable: {exc}") from None
        if not asset.object_sha256:
            raise ValueError(f"asset {asset_id} not found or has no stored object")
        data = await ctx.resolve("storage").get(object_key(asset.object_sha256))
        if data is None:
            raise ValueError(f"object bytes missing for asset {asset_id}")

        name = getattr(asset, "name", None) or "document"
        ext = Path(name).suffix.lower()
        mime = (getattr(asset, "mime_type", None) or "").lower()
        if ext in _IMAGE_EXTS or mime.startswith("image/"):
            return (
                f"'{name}' is an image. Call the `vision` tool with asset_id {asset_id} "
                "to analyze its visual content."
            )

        if pages is not None:
            if ext not in _PAGE_AXIS_EXTS:
                raise ValueError(
                    f"'{name}' has no page axis ({ext or 'unknown format'}): pages is only "
                    "supported for PDF and PowerPoint slide-based formats "
                    "(PPTX/POTX/PPSX). Omit pages to read the whole document, and never "
                    "substitute a full-document read for a page-scoped request."
                )
            if ext == ".pdf":
                text = await asyncio.to_thread(_extract_pdf_pages, data, name, pages)
            else:
                text = await asyncio.to_thread(_extract_pptx_slides, data, name, pages)
            if len(text) > MAX_OUTPUT_CHARS:
                total = len(text)
                text = text[:MAX_OUTPUT_CHARS] + f"\n\n[...truncated; {total} chars total]"
            return f"'{name}' — pages {','.join(str(p) for p in pages)}:\n\n" + text

        try:
            text = await extract_document_text(data, name, llm)
        except UnsupportedFileType as exc:
            return (
                f"Cannot extract text from '{name}': {exc} — you must NOT substitute "
                "another document from the conversation for this one; tell the user this "
                "file could not be read."
            )
        footer = ""
        if ext == ".doc":
            footer = await _mint_doc_images(data, asset, ctx)
        if not text.strip():
            return (
                f"'{name}' was parsed but contains no extractable text — it may be a "
                "scanned/image-only document — and you must NOT summarize a different "
                "document in its place." + footer
            )
        if len(text) > MAX_OUTPUT_CHARS:
            total = len(text)
            text = text[:MAX_OUTPUT_CHARS] + f"\n\n[...truncated; {total} chars total]"
        return text + footer

    runtime.register(
        define_tool(
            name="read_document",
            description="Extract the text content of a document the user attached to the "
            "chat or has open in the viewer (PDF, Word .doc/.docx, Excel .xlsx/.xlsm, "
            "PowerPoint .pptx/.potx/.ppsx, "
            "txt/markdown/csv/json, subtitles). Attachments arrive as "
            "[Attached: <filename> (asset_id <id>)]; the Viewer Access Context note hands "
            "you an asset_id the same way. Whenever the "
            "user asks about the content of such a document, you MUST call this tool "
            "with that asset_id instead of guessing or claiming the file cannot be read. "
            "Pass `pages` to read specific pages/slides (\"3\", \"2,5\", \"1-3\") — "
            "Only PDF and PowerPoint slide-based formats (PPTX/POTX/PPSX) support pages; "
            "for any other format omit the parameter (page-scoped requests are rejected, "
            "never answered with a full-document substitute). "
            "Embedded pictures in a legacy .doc come back with their own asset_ids at the "
            "end of the result — pass those ids to the `vision` tool to read the image or "
            "table. For attached images and screenshots use the `vision` tool instead.",
            parameters={
                "type": "object",
                "properties": {
                    "asset_id": {
                        "type": "string",
                        "description": "Cloud-drive asset id of the document "
                        "(from the [Attached: ...] or Viewer Access Context note).",
                    },
                    "pages": {
                        "type": "string",
                        "description": '1-based page/slide spec, e.g. "3", "2,5", "1-3". '
                        "Omit to read the whole document. Only PDF and PowerPoint "
                        "slide-based formats (PPTX/POTX/PPSX) support pages.",
                    },
                },
                "required": ["asset_id"],
            },
            output=ToolOutput(
                schema={"type": "string"}, render=lambda args, value: [text_block(value)]
            ),
            execute=read_document,
        )
    )
